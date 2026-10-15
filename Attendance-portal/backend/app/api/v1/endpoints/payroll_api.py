import csv
import io
import logging
from datetime import datetime
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Request, Response
from sqlalchemy import text
from sqlalchemy.orm import Session
from app.db.session import get_db
from app.services import payroll_service

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/payroll", tags=["Payroll"])


def _extract_admin_user(request: Request, db: Session):
    from app.core.dependencies import get_current_user
    from fastapi.security import HTTPAuthorizationCredentials

    auth_h = request.headers.get("authorization", "") if request else ""
    token = None
    if auth_h.startswith("Bearer "):
        token = auth_h.split(" ", 1)[1].strip()

    creds = HTTPAuthorizationCredentials(scheme="Bearer", credentials=token) if token else None
    try:
        user = get_current_user(request=request, credentials=creds, db=db)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=401, detail=f"Authentication failed: {str(exc)}")

    if not user:
        raise HTTPException(status_code=401, detail="Authenticated user was not found.")
    role = str(getattr(user, "role", "") or "").strip().lower()
    if role not in {"admin", "superadmin", "hr", "administrator", "manager"}:
        raise HTTPException(status_code=403, detail="Forbidden: Admin or manager access required for payroll operations.")
    return user


@router.get("")
def get_payroll_list(
    search: str = "",
    month: str = "",
    year: str = "",
    department: str = "",
    status: str = "",
    worker_type: str = "",
    has_exceptions: str = "",
    force_recalculate: bool = False,
    request: Request = None,
    db: Session = Depends(get_db)
):
    """
    Returns actual, database-backed payroll records for all employees.
    Calculates calendar days dynamically and provides reconciliation summary.
    """
    _extract_admin_user(request, db)
    return payroll_service.get_all_payroll_records(
        db=db,
        search=search,
        month=month,
        year=year,
        department=department,
        status=status,
        worker_type=worker_type,
        has_exceptions=has_exceptions,
        force_recalculate=force_recalculate
    )


@router.get("/settings")
def get_payroll_settings_api(request: Request, db: Session = Depends(get_db)):
    """Returns company-wide payroll rules and intern configuration."""
    _extract_admin_user(request, db)
    return payroll_service.get_payroll_settings(db)


@router.post("/settings")
@router.put("/settings")
async def update_payroll_settings_api(request: Request, db: Session = Depends(get_db)):
    """Updates company-wide payroll rules and intern configuration."""
    user = _extract_admin_user(request, db)
    body = await request.json()
    u_name = f"{user.first_name or ''} {user.last_name or ''}".strip() or "Admin"
    return payroll_service.update_payroll_settings(db, body, user_name=u_name)


@router.post("/calculate")
@router.post("/generate")
async def calculate_monthly_payroll_api(request: Request, db: Session = Depends(get_db)):
    """Recalculates dynamic payroll for the selected month and year."""
    _extract_admin_user(request, db)
    body = await request.json()
    month = body.get("month", "")
    year = str(body.get("year", ""))
    return payroll_service.get_all_payroll_records(
        db=db,
        month=month,
        year=year,
        force_recalculate=True
    )


@router.post("/finalize")
async def finalize_payroll_api(request: Request, db: Session = Depends(get_db)):
    """Finalizes payroll and creates official payslips for Employee Portal."""
    user = _extract_admin_user(request, db)
    body = await request.json()
    month = body.get("month", "")
    year = int(body.get("year", 2026))
    allow_with_exceptions = bool(body.get("allow_with_exceptions", False))
    u_name = f"{user.first_name or ''} {user.last_name or ''}".strip() or "Admin"
    try:
        return payroll_service.finalize_monthly_payroll(
            db, month=month, year=year, user_name=u_name, allow_with_exceptions=allow_with_exceptions
        )
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))


@router.post("/reopen")
async def reopen_payroll_api(request: Request, db: Session = Depends(get_db)):
    """Reopens a finalized payroll back to Draft status."""
    user = _extract_admin_user(request, db)
    body = await request.json()
    month = body.get("month", "")
    year = int(body.get("year", 2026))
    u_name = f"{user.first_name or ''} {user.last_name or ''}".strip() or "Admin"
    return payroll_service.reopen_monthly_payroll(db, month=month, year=year, user_name=u_name)


@router.get("/history")
def get_payroll_history_api(request: Request, db: Session = Depends(get_db)):
    """Returns all historical payroll runs."""
    _extract_admin_user(request, db)
    return payroll_service.get_payroll_history(db)


@router.post("/intern/convert")
async def intern_conversion_api(request: Request, db: Session = Depends(get_db)):
    """HR records intern conversion, extension, or status decision."""
    user = _extract_admin_user(request, db)
    body = await request.json()
    emp_code = body.get("emp_code")
    if not emp_code:
        raise HTTPException(status_code=400, detail="emp_code is required.")
    u_name = f"{user.first_name or ''} {user.last_name or ''}".strip() or "Admin"
    return payroll_service.record_intern_conversion(db, emp_code, body, user_name=u_name)


@router.put("/employee/{emp_code}/salary")
async def update_employee_salary_api(emp_code: str, request: Request, db: Session = Depends(get_db)):
    """Updates employee salary/stipend and PT configuration."""
    _extract_admin_user(request, db)
    body = await request.json()
    try:
        # Also update pt_applicable on profile_master if provided
        if "pt_applicable" in body or "ptApplicable" in body:
            pt_val = bool(body.get("pt_applicable") if "pt_applicable" in body else body.get("ptApplicable"))
            from sqlalchemy import text
            db.execute(text("UPDATE profile_master SET pt_applicable = :pt WHERE emp_code = :c"), {"pt": pt_val, "c": emp_code})
            db.execute(text("UPDATE payroll_profiles SET pt_applicable = :pt WHERE emp_code = :c"), {"pt": pt_val, "c": emp_code})
            db.commit()

        # Also update employment_type if provided
        if "worker_type" in body or "employment_type" in body:
            wt = body.get("worker_type") or body.get("employment_type")
            from sqlalchemy import text
            db.execute(text("UPDATE profile_master SET employment_type = :wt WHERE emp_code = :c"), {"wt": wt, "c": emp_code})
            db.commit()

        res = payroll_service.update_employee_salary_configuration(db, emp_code, body)
        return {"status": "success", "message": f"Salary updated for {emp_code}.", "data": res, **res}
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))



@router.post("/validate")
async def validate_payroll_records(request: Request, db: Session = Depends(get_db)):
    _extract_admin_user(request, db)
    body = await request.json()
    records = body.get("records", [])
    valid = []
    invalid = []
    for idx, r in enumerate(records, start=1):
        if not r.get("employeeId") and not r.get("employee_id") and not r.get("emp_code"):
            invalid.append({"row": idx, "errors": ["Missing Employee ID"]})
        else:
            valid.append({"row": idx, **r, "errors": []})
    return {"valid": valid, "invalid": invalid}


@router.post("/import")
async def import_payroll_records(request: Request, db: Session = Depends(get_db)):
    _extract_admin_user(request, db)
    body = await request.json()
    records = body.get("records", [])
    count = 0
    for r in records:
        emp_c = r.get("employeeId") or r.get("employee_id") or r.get("emp_code")
        if not emp_c:
            continue
        m = r.get("month") or "September"
        y = int(r.get("year") or 2026)
        try:
            payroll_service.generate_payslip_for_employee_record(db, emp_c, m, y)
            count += 1
        except Exception:
            pass
    return {"imported": count, "message": f"Successfully imported {count} payroll records."}


@router.post("/{record_id}/generate-payslip")
async def generate_payslip_api(record_id: str, request: Request, db: Session = Depends(get_db)):
    """
    Generates a single real payslip for the specified employee record.
    Prevents duplicates and strictly verifies salary is configured.
    """
    _extract_admin_user(request, db)

    # Parse record_id: e.g. "12", "rec-LE724-September-2026", or "LE724"
    emp_code = None
    month = "September"
    year = 2026

    if str(record_id).startswith("rec-"):
        parts = record_id.split("-")
        if len(parts) >= 4:
            emp_code = parts[1]
            month = parts[2]
            year = int(parts[3])
    elif not str(record_id).isdigit():
        emp_code = record_id
    else:
        # Check if record_id is an existing payslip id
        from sqlalchemy import text
        row = db.execute(text("SELECT emp_code, month, year FROM payslips WHERE id = :id"), {"id": int(record_id)}).mappings().first()
        if row:
            emp_code = row["emp_code"]
            month = row["month"]
            year = row["year"]

    # Also check request body for overrides
    try:
        body = await request.json()
        if body.get("emp_code"):
            emp_code = body.get("emp_code")
        if body.get("month"):
            month = body.get("month")
        if body.get("year"):
            year = int(body.get("year"))
    except Exception:
        pass

    if not emp_code:
        raise HTTPException(status_code=400, detail="Employee ID could not be identified from record.")

    try:
        payslip = payroll_service.generate_payslip_for_employee_record(db, emp_code, month, year)
        return {
            "status": "success",
            "message": f"Payslip generated successfully for {emp_code} ({month} {year}).",
            "data": payslip,
            **payslip
        }
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/{record_id}/approve")
def approve_payroll_api(record_id: str, request: Request, db: Session = Depends(get_db)):
    _extract_admin_user(request, db)
    try:
        res = payroll_service.approve_payroll_record(db, record_id)
        return {"status": "success", "message": "Payroll approved.", "data": res, **res}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.put("/{record_id}")
async def update_payroll_api(record_id: str, request: Request, db: Session = Depends(get_db)):
    _extract_admin_user(request, db)
    body = await request.json()

    # Determine emp_code
    emp_code = record_id
    if str(record_id).startswith("rec-"):
        parts = record_id.split("-")
        if len(parts) >= 2:
            emp_code = parts[1]
    elif str(record_id).isdigit():
        from sqlalchemy import text
        row = db.execute(text("SELECT emp_code FROM payslips WHERE id = :id"), {"id": int(record_id)}).mappings().first()
        if row:
            emp_code = row["emp_code"]

    emp_code = body.get("emp_code") or body.get("employeeId") or emp_code

    try:
        res = payroll_service.update_employee_salary_configuration(db, emp_code, body)
        return {"status": "success", "message": f"Salary updated for {emp_code}.", "data": res, **res}
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/{record_id}/payslip/download")
def download_payslip_api(
    record_id: str,
    month: str = "",
    year: str = "",
    request: Request = None,
    db: Session = Depends(get_db)
):
    _extract_admin_user(request, db)
    try:
        clean_code = record_id
        target_month = month
        target_year = year

        if record_id.startswith("rec-"):
            parts = record_id.split("-")
            if len(parts) >= 2:
                clean_code = parts[1]
            if len(parts) >= 3 and not target_month:
                target_month = parts[2]
            if len(parts) >= 4 and not target_year:
                target_year = parts[3]
        elif record_id.isdigit():
            num_id = int(record_id)
            match = db.execute(
                text("SELECT emp_code, month, year FROM payroll_records WHERE id = :id"),
                {"id": num_id}
            ).mappings().first()
            if not match:
                match = db.execute(
                    text("SELECT emp_code, month, year FROM payslips WHERE id = :id"),
                    {"id": num_id}
                ).mappings().first()
            if match:
                clean_code = match["emp_code"]
                if not target_month:
                    target_month = match["month"]
                if not target_year:
                    target_year = str(match["year"])

        now = datetime.now()
        if not target_month:
            target_month = "October" if now.year == 2026 else payroll_service.MONTH_NAMES[now.month - 1]
        if not target_year:
            target_year = "2026" if now.year == 2026 else str(now.year)

        pdf_bytes = payroll_service.generate_payslip_pdf(
            db=db,
            payslip_id_or_code=record_id,
            month=month,
            year=year
        )
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename=Payslip_{clean_code}_{target_month}_{target_year}.pdf"
            }
        )
    except ValueError as ve:
        raise HTTPException(status_code=404, detail=str(ve))
    except Exception as e:
        logger.error(f"Payslip PDF download failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="Unable to generate payslip. Please try again.")


@router.get("/employee/{emp_code}/breakdown")
def get_employee_breakdown_api(
    emp_code: str,
    month: str = "",
    year: str = "",
    request: Request = None,
    db: Session = Depends(get_db)
):
    """Returns granular calculation breakdown for employee modal."""
    _extract_admin_user(request, db)
    data = payroll_service.get_all_payroll_records(
        db=db, search=emp_code, month=month, year=year
    )
    rec = next((r for r in data.get("records", []) if r.get("employeeId") == emp_code or r.get("emp_code") == emp_code), None)
    if not rec:
        raise HTTPException(status_code=404, detail=f"Employee {emp_code} not found in payroll.")
    return {"status": "success", "data": rec, **rec}


@router.get("/export")
def export_payroll_api(
    format: str = "csv",
    search: str = "",
    month: str = "",
    year: str = "",
    department: str = "",
    status: str = "",
    worker_type: str = "",
    has_exceptions: str = "",
    request: Request = None,
    db: Session = Depends(get_db)
):
    """Exports comprehensive payroll report in Excel (.xlsx) or CSV format matching Section 57."""
    _extract_admin_user(request, db)
    data = payroll_service.get_all_payroll_records(
        db=db, search=search, month=month, year=year,
        department=department, status=status, worker_type=worker_type,
        has_exceptions=has_exceptions
    )
    records = data.get("records", [])

    headers = [
        "Sr.no", "Employee ID", "Employee Name", "Designation", "Department", "Worker Type",
        "Gross Salary", "Days in month", "Paid Days", "Leave Taken",
        "Fixed Gross Salary", "Basic + DA", "HRA", "Flexi Allowance", "Gross",
        "Earn basic", "Earn HRA", "Earn Other Allo", "Flexi Allowance2", "Earn Gross",
        "PF WAGES", "PF", "ESIC", "PT", "LWF", "TDS",
        "Leave Deduction", "Other Deduction", "Arrears",
        "Net Salary", "Payroll Status", "Exceptions"
    ]

    def format_row(idx, r):
        exc_str = "; ".join([e.get("message", "") for e in r.get("exceptions", [])]) if r.get("exceptions") else "None"
        gross_val = r.get("monthlyCompensation", 0.0)
        cal_d = r.get("calendarDays", 30)
        paid_d = r.get("paidEquivalentDays", 0.0)
        leave_taken = r.get("leaveTaken", 0.0)

        # Fixed structure
        fixed_gross = r.get("fixedGrossSalary", gross_val)
        basic_val = r.get("basicSalary", round(gross_val * 0.5, 2))
        hra_val = r.get("hra", round(gross_val * 0.25, 2))
        flexi_val = r.get("flexiAllowance", round(gross_val - basic_val - hra_val, 2))

        # Earned components
        earn_basic = r.get("earnBasic", round(basic_val * (paid_d / cal_d), 2) if cal_d else 0.0)
        earn_hra = r.get("earnHra", round(hra_val * (paid_d / cal_d), 2) if cal_d else 0.0)
        earn_other = r.get("earnOtherAllo", round(flexi_val * (paid_d / cal_d), 2) if cal_d else 0.0)
        earn_gross = r.get("earnGross", r.get("grossEarnedSalary", 0.0))

        # Deductions
        pf_wages = float(r.get("pfWages") or 0.0)
        pf = float(r.get("pfDeduction") or 0.0)
        esic = float(r.get("esicDeduction") or 0.0)
        pt = float(r.get("ptDeducted") or 0.0)
        lwf = float(r.get("lwfDeduction") or 0.0)
        tds = float(r.get("tdsDeduction") or 0.0)
        leave_ded = float(r.get("lopDeduction") or 0.0)
        other_ded = float(r.get("otherDeductions") or 0.0)
        arrears = float(r.get("arrears") or 0.0)
        net_sal = float(r.get("netSalary") or 0.0)

        return [
            idx,
            r.get("employeeId"),
            r.get("employeeName"),
            r.get("designation"),
            r.get("department"),
            r.get("workerType"),
            gross_val,
            cal_d,
            paid_d,
            leave_taken,
            fixed_gross,
            basic_val,
            hra_val,
            flexi_val,
            fixed_gross,
            earn_basic,
            earn_hra,
            earn_other,
            0.0,
            earn_gross,
            pf_wages,
            pf,
            esic,
            pt,
            lwf,
            tds,
            leave_ded,
            other_ded,
            arrears,
            net_sal,
            r.get("status"),
            exc_str
        ]

    if format.lower() == "xlsx":
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Payroll_{month or 'Month'}"

        ws.append(headers)

        # Style Header Row matching Company Reference Sheet
        header_font = Font(name="Calibri", size=11, bold=True, color="000000")
        yellow_fill = PatternFill(start_color="FFF200", end_color="FFF200", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        default_fill = PatternFill(start_color="EAECEF", end_color="EAECEF", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D0D5DD'),
            right=Side(style='thin', color='D0D5DD'),
            top=Side(style='thin', color='D0D5DD'),
            bottom=Side(style='thin', color='D0D5DD')
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            
            # Yellow group (Columns 11 to 15: Fixed Gross Salary, Basic + DA, HRA, Flexi Allowance, Gross)
            if 11 <= col_idx <= 15:
                cell.fill = yellow_fill
            # Green group (Columns 8 to 10: Days in month, Paid Days, Leave Taken, and Columns 27-29: Leave Ded, Other Ded, Arrears)
            elif 8 <= col_idx <= 10 or 27 <= col_idx <= 29:
                cell.fill = green_fill
            else:
                cell.fill = default_fill

        ws.row_dimensions[1].height = 28

        for idx, r in enumerate(records, start=1):
            row_data = format_row(idx, r)
            ws.append(row_data)
            current_row = ws.max_row
            ws.row_dimensions[current_row].height = 20
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=current_row, column=col_idx)
                c.border = thin_border
                c.font = Font(name="Calibri", size=10)
                # Numeric alignments
                if isinstance(c.value, (int, float)):
                    c.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return Response(
            content=stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=Payroll_{month or 'All'}_{year or '2026'}.xlsx"}
        )
    else:
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(headers)
        for idx, r in enumerate(records, start=1):
            writer.writerow(format_row(idx, r))
        return Response(
            content=stream.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=Payroll_{month or 'All'}_{year or '2026'}.csv"}
        )

